import base64
import re
from pathlib import Path
from io import BytesIO
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn


# ---------- DOCX -> HTML (for loading into the editor) ----------

def get_docx_as_html(file_path, file_name):
    target_path = Path(file_path) / file_name
    doc = Document(target_path)

    image_rels = {
        rel_id: rel.target_part
        for rel_id, rel in doc.part.rels.items()
        if "image" in rel.reltype
    }

    html_parts = []

    def run_to_html(run):
        text = run.text.replace("<", "&lt;").replace(">", "&gt;")
        if not text:
            return ""
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.underline:
            text = f"<u>{text}</u>"
        return text

    def paragraph_to_html(paragraph):
        style_name = (paragraph.style.name or "").lower()
        tag = "p"
        if "heading" in style_name:
            level = "".join(filter(str.isdigit, style_name)) or "1"
            tag = f"h{min(int(level), 6)}"

        inner = ""
        for el in paragraph._element.iter():
            if el.tag == qn("a:blip"):
                rel_id = el.get(qn("r:embed"))
                if rel_id in image_rels:
                    part = image_rels[rel_id]
                    mime = part.content_type
                    b64 = base64.b64encode(part.blob).decode("utf-8")
                    inner += f'<img src="data:{mime};base64,{b64}" style="max-width:100%;">'

        for run in paragraph.runs:
            inner += run_to_html(run)

        if not inner.strip():
            inner = "<br>"
        return f"<{tag}>{inner}</{tag}>"

    def table_to_html(table):
        rows_html = ""
        for row in table.rows:
            cells_html = "".join(f"<td>{cell.text}</td>" for cell in row.cells)
            rows_html += f"<tr>{cells_html}</tr>"
        return f'<table border="1" style="border-collapse:collapse;width:100%;">{rows_html}</table>'

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            html_parts.append(paragraph_to_html(Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            from docx.table import Table
            html_parts.append(table_to_html(Table(child, doc)))

    return "\n".join(html_parts)


# ---------- HTML -> DOCX (for saving edits) ----------

def _add_runs_from_node(paragraph, node, bold=False, italic=False, underline=False):
    """Recursively walk inline HTML nodes, adding runs with correct formatting."""
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            run.bold = bold
            run.italic = italic
            run.underline = underline
        return

    name = node.name
    if name in ("strong", "b"):
        bold = True
    elif name in ("em", "i"):
        italic = True
    elif name == "u":
        underline = True
    elif name == "br":
        paragraph.add_run().add_break()
        return
    elif name == "img":
        src = node.get("src", "")
        m = re.match(r"data:(image/\w+);base64,(.+)", src)
        if m:
            img_data = base64.b64decode(m.group(2))
            paragraph.add_run().add_picture(BytesIO(img_data), width=Inches(4))
        return

    for child in node.children:
        _add_runs_from_node(paragraph, child, bold, italic, underline)


def save_html_as_docx(file_path, file_name, html_content):
    """Rebuilds the docx from scratch based on the edited HTML."""
    target_path = Path(file_path) / file_name
    doc = Document()  # fresh document — replaces old content entirely

    soup = BeautifulSoup(html_content, "html.parser")

    heading_map = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}

    for element in soup.find_all(["p", "h1", "h2", "h3", "h4", "h5", "h6", "table"], recursive=False):
        if element.name == "table":
            rows = element.find_all("tr")
            if not rows:
                continue
            n_cols = len(rows[0].find_all("td"))
            table = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                cells = row.find_all("td")
                for c_idx, cell in enumerate(cells):
                    table.cell(r_idx, c_idx).text = cell.get_text()
        else:
            if element.name in heading_map:
                paragraph = doc.add_paragraph(style=f"Heading {heading_map[element.name]}")
            else:
                paragraph = doc.add_paragraph()
            for child in element.children:
                _add_runs_from_node(paragraph, child)

    doc.save(target_path)
    return str(target_path)